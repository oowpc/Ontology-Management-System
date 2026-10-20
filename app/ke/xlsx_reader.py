"""xlsx 读取（openpyxl）： bytes → 首个工作表按表头的行 dict 列表。"""
from __future__ import annotations

import io

from openpyxl import load_workbook


def read_xlsx(raw: bytes) -> tuple[str, list[dict]]:
    """读第一个工作表：返回 (工作表名, 行 dict 列表)（单元格 None→''，全部转 str 去首尾空白）。"""
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"xlsx 解析失败：{e}") from e
    ws = wb.worksheets[0]
    # 协作系统/WPS 导出的表 dimension 常不准（声明只有表头行）；read_only 下强制重扫
    try:
        ws.reset_dimensions()
    except AttributeError:
        pass  # 非 read_only 模式无需
    rows = list(ws.iter_rows(values_only=True))
    sheet_name = ws.title or ""
    wb.close()
    if not rows:
        raise ValueError("xlsx 第一个工作表为空")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        row = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            v = r[i] if i < len(r) else None
            row[h] = "" if v is None else str(v).strip()
        if any(row.values()):  # 全空行跳过
            out.append(row)
    return sheet_name, out
